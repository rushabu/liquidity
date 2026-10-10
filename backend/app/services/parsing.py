"""File -> raw rows. No type inference and no validation: cells come out as the
reader found them (datetime / int / float / str / None). pandas is not used.

Row numbers are 1-based as the user sees them in Excel/CSV: header = row 1,
first data row = row 2. Completely blank rows are skipped but still counted
in the numbering, so reported rows always match the user's file."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date
from typing import Iterator

from openpyxl import load_workbook

from .errors import ApiError

SUPPORTED_EXT = (".xlsx", ".csv")


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[tuple[int, dict[str, object]]]    # (source_row, {header: raw value})


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _clean_header(v) -> str:
    return "" if v is None else str(v).strip().lstrip("﻿")


def _iter_xlsx(content: bytes) -> Iterator[tuple]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()       # some writers store a wrong <dimension>; read everything
        yield from ws.iter_rows(values_only=True)
        wb.close()
    except ApiError:
        raise
    except Exception as e:
        raise ApiError(422, "FILE_UNREADABLE",
                       f"The Excel file could not be read ({e.__class__.__name__}). Is it a valid .xlsx?") from e


def _iter_csv(content: bytes) -> Iterator[tuple]:
    try:
        text = content.decode("utf-8-sig")      # BOM tolerant
    except UnicodeDecodeError as e:
        raise ApiError(422, "FILE_UNREADABLE", "The CSV file is not valid UTF-8.") from e
    yield from csv.reader(io.StringIO(text, newline=""))


def parse_file(content: bytes, filename: str) -> ParsedFile:
    name = (filename or "").lower()
    if not name.endswith(SUPPORTED_EXT):
        raise ApiError(422, "FILE_TYPE_UNSUPPORTED",
                       f"Unsupported file type. Upload one of: {', '.join(SUPPORTED_EXT)}",
                       {"filename": filename})
    if not content:
        raise ApiError(422, "FILE_EMPTY", "The uploaded file is empty.")

    it = _iter_xlsx(content) if name.endswith(".xlsx") else _iter_csv(content)
    headers: list[str] | None = None
    rows: list[tuple[int, dict[str, object]]] = []
    for n, raw in enumerate(it, start=1):
        if headers is None:
            headers = [_clean_header(c) for c in raw]
            continue
        if all(_blank(c) for c in raw):
            continue
        cells = list(raw) + [None] * (len(headers) - len(raw))
        rows.append((n, {h: (c.strip() if isinstance(c, str) else c)
                         for h, c in zip(headers, cells) if h}))
    if headers is None or not rows:
        raise ApiError(422, "FILE_EMPTY", "The file has no data rows.")
    return ParsedFile(headers=[h for h in headers if h], rows=rows)


_AS_OF = re.compile(r"(?<!\d)(\d{2})(\d{2})(\d{4})(?!\d)")


def as_of_from_filename(filename: str) -> date | None:
    """'asset_book - 10102026.xlsx' -> 2026-10-10 (DDMMYYYY). None if absent/invalid."""
    for m in _AS_OF.finditer(filename or ""):
        try:
            return date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            continue
    return None
